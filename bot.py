import html
import io
import logging
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from telegram import Update
from telegram.constants import ParseMode
from telegram.ext import (Application, CommandHandler, ContextTypes,
                          MessageHandler, filters)

import rasch

logging.basicConfig(level=logging.INFO)

BOT_TOKEN = os.environ["BOT_TOKEN"]
ADMIN_ID = int(os.environ["ADMIN_ID"]) if os.environ.get("ADMIN_ID") else None  # ixtiyoriy
SKIP_HEADERS = {"jami", "umumiy", "total", "sum", "ball", "foiz", "%", "natija"}
FONT = "Arial"

HELP = (
    "📊 <b>Baholash boti (Rasch modeli, Milliy sertifikat)</b>\n\n"
    "Excel (.xlsx) faylni yuboring:\n"
    "• 1-ustun — o'quvchi ismi\n"
    "• keyingi ustunlar — savollar (1, 2, 3 ...)\n"
    "• har bir katakda: to'g'ri = 1, xato = 0 (bo'sh katak bo'lmasin)\n"
    "• 'Jami', 'Ball' kabi ustunlar bo'lsa, ular o'tkazib yuboriladi\n\n"
    "Bot har bir o'quvchiga ball (0–75), foiz va daraja qo'yib, jadval qaytaradi."
)


# ---------------------------------------------------------------- Excel o'qish
def _to01(v):
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)) and v in (0, 1):
        return int(v)
    if isinstance(v, str) and v.strip() in ("0", "1"):
        return int(v.strip())
    return None


def parse_excel(data: bytes):
    ws = load_workbook(io.BytesIO(data), data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise ValueError("Jadval bo'sh: sarlavha va kamida bitta o'quvchi qatori kerak.")
    header = rows[0]
    filled = [i for i, h in enumerate(header) if h not in (None, "")]
    if len(filled) < 2:
        raise ValueError("Sarlavha qatorida savollar ustunlari topilmadi.")
    last = max(filled)
    q_cols = [c for c in range(1, last + 1)
              if str(header[c] or "").strip().lower() not in SKIP_HEADERS]
    q_names = [str(header[c]) if header[c] not in (None, "") else f"Savol {c}"
               for c in q_cols]

    names, matrix, errors = [], [], []
    for r_idx, row in enumerate(rows[1:], start=2):
        if all(c in (None, "") for c in row):
            continue
        name = str(row[0]).strip() if row[0] not in (None, "") else f"Noma'lum (qator {r_idx})"
        vals = []
        for c in q_cols:
            v = _to01(row[c] if c < len(row) else None)
            if v is None:
                errors.append(f"{get_column_letter(c + 1)}{r_idx}")
            vals.append(v)
        names.append(name)
        matrix.append(vals)

    if errors:
        shown = ", ".join(errors[:15]) + (f" ... (jami {len(errors)} ta)" if len(errors) > 15 else "")
        raise ValueError("Faqat 1 yoki 0 bo'lishi kerak. Noto'g'ri/bo'sh kataklar: " + shown)
    if len(names) < 2 or len(q_names) < 2:
        raise ValueError("Kamida 2 ta o'quvchi va 2 ta savol kerak.")
    return names, q_names, matrix


# ---------------------------------------------------------------- Excel yozish
def _style_header(ws, ncols):
    fill = PatternFill("solid", start_color="1F4E78")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _finish(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
    ws.freeze_panes = "A2"


def build_report(names, q_names, res) -> bytes:
    n = len(names)
    order = sorted(range(n), key=lambda i: (-res["ball"][i], names[i].lower()))
    wb = Workbook()

    ws = wb.active
    ws.title = "Natijalar"
    ws.append(["O'rin", "Ism", "To'g'ri javoblar", "Jami savol", "Qobiliyat θ (logit)",
               "Xato (SE)", "Ball (0–75)", "Foiz (%)", "Daraja"])
    for i in order:
        rank = 1 + sum(1 for j in range(n) if res["ball"][j] > res["ball"][i])
        ws.append([rank, names[i], int(res["raw_total"][i]), res["n_items"],
                   round(float(res["theta"][i]), 3), round(float(res["se"][i]), 3),
                   float(res["ball"][i]), float(res["foiz"][i]), res["daraja"][i]])
    _style_header(ws, 9)
    _finish(ws, [8, 28, 14, 12, 18, 12, 12, 10, 18])

    ws2 = wb.create_sheet("Savollar")
    ws2.append(["Savol", "To'g'ri javob soni", "To'g'ri (%)", "Qiyinlik b (logit)",
                "Xato (SE)", "Infit MNSQ", "Outfit MNSQ", "Holat"])
    for it in res["items"]:
        ws2.append([q_names[it["index"]], it["correct"], it["pct"],
                    None if it["b"] is None else round(it["b"], 3),
                    None if it["se"] is None else round(it["se"], 3),
                    None if it["infit"] is None else round(it["infit"], 2),
                    None if it["outfit"] is None else round(it["outfit"], 2),
                    it["status"]])
    _style_header(ws2, 8)
    _finish(ws2, [14, 18, 12, 18, 12, 12, 12, 36])

    ws3 = wb.create_sheet("Metodika")
    rel = res["reliability"]
    lines = [
        "Hisoblash metodikasi (shaffoflik uchun)",
        "",
        "Model: dixotomik Rasch modeli, P(to'g'ri) = 1 / (1 + exp(-(θ - b))).",
        "Baholash usuli: JMLE (Newton-Raphson) + (k-1)/k siljish tuzatishi; o'rtacha qiyinlik = 0 logit.",
        "Hamma to'g'ri / hamma xato savollar va 0 yoki 100% olgan o'quvchilar kalibrovkaga kirmaydi.",
        f"0 yoki to'liq ball olgan o'quvchilar uchun {rasch.EXTREME_CORR} tuzatish ishlatilgan.",
        f"θ → ball: ball = {rasch.MAX_BALL / 2} + {rasch.MAX_BALL / (2 * rasch.LOGIT_RANGE):.3f} × θ, "
        f"0..{int(rasch.MAX_BALL)} oralig'ida cheklangan (θ = ±{rasch.LOGIT_RANGE:g} logit ↔ 0 va {int(rasch.MAX_BALL)}).",
        "Foiz = ball / 75 × 100.",
        "Daraja: >70 A+; 65–69.9 A; 60–64.9 B+; 55–59.9 B; 50–54.9 C+; 46–49.9 C; <46 sertifikat yo'q.",
        "",
        f"O'quvchilar: {res['n_persons']} (kalibrovkada: {res['n_persons_calibrated']})",
        f"Savollar: {res['n_items']} (hisobga olingan: {res['n_items_used']})",
        "Ishonchlilik (person reliability): " + ("hisoblanmadi" if rel is None else f"{rel:.2f}"),
        "Konvergensiya: " + ("ha" if res["converged"] else "YO'Q — natijalarga ehtiyot bo'ling"),
        "",
        "ESLATMA: DTM ning θ → ball o'tkazish formulasi ommaga e'lon qilinmagan, shuning uchun "
        "ballar rasmiy Milliy sertifikat natijasi bilan aynan mos kelmasligi mumkin — bu mashq/tayyorgarlik "
        "uchun ochiq va tekshiriladigan baho. Rasch modelida bir xil xom ball (hisobga olingan savollar "
        "bo'yicha) bir xil ball beradi; farq savollarning qiyinligi hisobga olinishida.",
    ]
    if res["n_persons"] < 30:
        lines.append("DIQQAT: o'quvchilar soni 30 tadan kam — Rasch baholari taxminiy va beqaror bo'lishi mumkin.")
    for ln in lines:
        ws3.append([ln])
    ws3.column_dimensions["A"].width = 120
    for row in ws3.iter_rows():
        for cell in row:
            cell.font = Font(name=FONT, bold=(cell.row == 1))
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------- chatdagi jadval
def chat_tables(names, res):
    n = len(names)
    order = sorted(range(n), key=lambda i: (-res["ball"][i], names[i].lower()))
    head = f"{'№':>3} {'Ism':<18} {'T/J':>7} {'Ball':>5} {'%':>6} Daraja"
    lines = []
    for pos, i in enumerate(order, start=1):
        nm = names[i][:18]
        lines.append(f"{pos:>3} {nm:<18} {int(res['raw_total'][i]):>3}/{res['n_items']:<3} "
                     f"{res['ball'][i]:>5.1f} {res['foiz'][i]:>5.1f}% {res['daraja'][i]}")
    chunks, cur = [], [head]
    size = len(head)
    for ln in lines:
        if size + len(ln) > 3300:
            chunks.append("\n".join(cur))
            cur, size = [head], len(head)
        cur.append(ln)
        size += len(ln) + 1
    chunks.append("\n".join(cur))
    return [f"<pre>{html.escape(c)}</pre>" for c in chunks]


# ---------------------------------------------------------------- handlerlar
def allowed(update: Update) -> bool:
    return ADMIN_ID is None or (update.effective_user is not None
                                and update.effective_user.id == ADMIN_ID)


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if allowed(update):
        await update.message.reply_text(HELP, parse_mode=ParseMode.HTML)


async def process_results(update: Update, context: ContextTypes.DEFAULT_TYPE):
    msg = update.message
    wait = await msg.reply_text("⏳ Hisoblanmoqda...")
    try:
        tg_file = await msg.document.get_file()
        data = bytes(await tg_file.download_as_bytearray())
        names, q_names, matrix = parse_excel(data)
        res = rasch.fit(matrix)
        report = build_report(names, q_names, res)
    except ValueError as e:
        await wait.edit_text(f"❌ {e}")
        return
    except Exception:
        logging.exception("Excel qayta ishlashda xato")
        await wait.edit_text("❌ Faylni o'qib bo'lmadi. .xlsx formatda ekanini tekshiring.")
        return

    await wait.delete()
    warn = []
    if res["n_persons"] < 30:
        warn.append("⚠️ O'quvchilar 30 tadan kam — Rasch baholari taxminiy.")
    if res["n_items_used"] < res["n_items"]:
        warn.append(f"ℹ️ {res['n_items'] - res['n_items_used']} ta savol (hamma to'g'ri/hamma xato) "
                    "hisobga olinmadi.")
    if not res["converged"]:
        warn.append("⚠️ Hisoblash to'liq yaqinlashmadi — natijalarga ehtiyot bo'ling.")
    rel = res["reliability"]
    summary = (f"✅ {res['n_persons']} ta o'quvchi, {res['n_items']} ta savol"
               + (f", ishonchlilik: {rel:.2f}" if rel is not None else "")
               + ("\n" + "\n".join(warn) if warn else ""))
    await msg.reply_text(summary)
    for t in chat_tables(names, res):
        await msg.reply_text(t, parse_mode=ParseMode.HTML)
    await msg.reply_document(document=io.BytesIO(report), filename="natijalar.xlsx",
                             caption="To'liq hisobot: natijalar, savollar tahlili va metodika")


async def on_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not allowed(update):
        return
    doc = update.message.document
    if (doc.file_name or "").lower().endswith(".xlsx"):
        await process_results(update, context)
    else:
        await update.message.reply_text("Faqat .xlsx (Excel) fayl yuboring.")


async def on_other(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if allowed(update):
        await update.message.reply_text(HELP, parse_mode=ParseMode.HTML)


def main():
    app = Application.builder().token(BOT_TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(MessageHandler(filters.Document.ALL, on_document))
    app.add_handler(MessageHandler(~filters.COMMAND & ~filters.Document.ALL, on_other))
    app.run_polling()


if __name__ == "__main__":
    main()
