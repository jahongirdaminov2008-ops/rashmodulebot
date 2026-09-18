import openpyxl


def export_submissions_to_excel(path: str, test_id: str, submissions, open_count: int, closed_count: int):
    """O'qituvchi tekshirishi uchun excel yaratadi: har savolga javob + bo'sh ball ustunlari (1/0 to'ldiriladi)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = test_id[:31]

    headers = ["student_id", "full_name"]
    headers += [f"O{i+1}_javob" for i in range(open_count)]
    headers += [f"Y{i+1}_javob" for i in range(closed_count)]
    headers += [f"O{i+1}_ball" for i in range(open_count)]
    headers += [f"Y{i+1}_ball" for i in range(closed_count)]
    ws.append(headers)

    for row in submissions:
        open_ans = (row["open_answers"] or "").split("|") if row["open_answers"] else []
        closed_ans = (row["closed_answers"] or "").split("|") if row["closed_answers"] else []
        open_ans += [""] * (open_count - len(open_ans))
        closed_ans += [""] * (closed_count - len(closed_ans))
        data_row = [row["student_id"], row["full_name"]] + open_ans[:open_count] + closed_ans[:closed_count]
        data_row += [""] * (open_count + closed_count)  # o'qituvchi 1/0 yozadigan bo'sh ustunlar
        ws.append(data_row)

    wb.save(path)


def read_graded_excel(path: str, open_count: int, closed_count: int):
    """O'qituvchi to'ldirgan (1/0 baholangan) excelni o'qiydi."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header, data_rows = rows[0], rows[1:]

    n_items = open_count + closed_count
    grade_start = len(header) - n_items

    results = []
    for row in data_rows:
        if not row or not row[0]:
            continue
        student_id, full_name = str(row[0]), row[1]
        grades_raw = row[grade_start:grade_start + n_items]
        grades = [int(g) if g is not None and str(g).strip() != "" else 0 for g in grades_raw]
        results.append((student_id, full_name, grades))
    return results


def build_results_excel(path: str, results):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Natijalar"
    ws.append(["ID", "F.I.Sh.", "To'g'ri javob", "Jami savol", "Foiz (%)", "Daraja"])
    for r in results:
        ws.append([r["student_id"], r["full_name"], r["raw_score"], r["max_score"],
                   round(r["percent"], 2), r["daraja"]])
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = width
    wb.save(path)
